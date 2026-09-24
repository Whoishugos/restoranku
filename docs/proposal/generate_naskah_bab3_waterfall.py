#!/usr/bin/env python3
"""Generate Bab 3 revision Word documents aligned with Gambar 3.1 Waterfall.

Creates:
  - Naskah_Perbaikan_Bab3_Waterfall.docx
  - Proposal-Revisi-22-September-2026.docx (Bab 3 excerpt for paste-in)

Topic: Restoran Kekupu Villa Jembrana — Laravel 12, QR Code meja, Midtrans,
Black-Box Testing & UAT. Not Ari Nata / IPO flowchart.
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT_DIR = Path("/workspace/docs/proposal")
ARTIFACTS = Path("/opt/cursor/artifacts")
PNG = OUT_DIR / "gambar-3-1-alur-waterfall.png"

FONT = "Times New Roman"


def _set_run_font(run, *, size=12, bold=False, italic=False, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _set_paragraph_format(p, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=True, space_after=6):
    fmt = p.paragraph_format
    fmt.alignment = align
    fmt.space_after = Pt(space_after)
    fmt.space_before = Pt(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first_line:
        fmt.first_line_indent = Cm(1.25)
    else:
        fmt.first_line_indent = Cm(0)


def add_heading_tnr(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        _set_run_font(run, size=14 if level == 1 else 12, bold=True)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    return h


def add_body(doc, text, *, first_line=True, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run_font(run, size=12, bold=bold, italic=italic)
    _set_paragraph_format(p, align=align, first_line=first_line)
    return p


def add_caption(doc, text, *, italic=False, bold=True):
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run_font(run, size=11, bold=bold, italic=italic)
    _set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, space_after=2)
    return p


def setup_page(doc):
    for section in doc.sections:
        section.top_margin = Cm(3)
        section.bottom_margin = Cm(3)
        section.left_margin = Cm(4)
        section.right_margin = Cm(3)
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def insert_gambar_3_1(doc):
    if not PNG.exists():
        raise FileNotFoundError(f"Missing flowchart image: {PNG}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run()
    run.add_picture(str(PNG), width=Cm(12.5))
    add_caption(doc, "Gambar 3.1 Alur Penelitian menggunakan metode waterfall", bold=True, italic=False)
    add_caption(
        doc,
        "Sumber: Diadaptasi dari model Waterfall (Pressman, 2015)",
        bold=False,
        italic=True,
    )


# --- Shared body paragraphs for 3.4 ---

INTRO_REPLACEMENT = (
    "Berdasarkan hasil revisi pembimbing, diagram alur penelitian pada Bab III yang "
    "sebelumnya menggunakan kerangka Input–Proses–Output (IPO) diganti dengan alur "
    "penelitian berbasis model Waterfall. Penggantian ini dilakukan agar tahapan "
    "penelitian selaras dengan proses rekayasa perangkat lunak yang digunakan untuk "
    "membangun sistem pemesanan Restoran Kekupu Villa Jembrana berbasis Laravel 12, "
    "pemindaian QR Code meja, serta integrasi pembayaran Midtrans. Alur penelitian "
    "yang diperbaiki digambarkan pada Gambar 3.1."
)

PROSEDUR_INTRO = (
    "Prosedur penelitian dalam pengembangan sistem informasi pemesanan Restoran "
    "Kekupu Villa Jembrana mengikuti metode Waterfall yang diadaptasi dari Pressman "
    "(2015). Setiap tahap diselesaikan secara berurutan sebagaimana ditunjukkan pada "
    "Gambar 3.1, dengan penjelasan sebagai berikut."
)

STAGES = [
    (
        "1. Analisis Kebutuhan",
        "Tahap awal penelitian dimulai setelah Mulai, yaitu analisis kebutuhan sistem "
        "pemesanan restoran. Peneliti merumuskan kebutuhan fungsional dan nonfungsional "
        "untuk pelanggan, staf dapur/pelayan, serta admin Restoran Kekupu Villa, termasuk "
        "kebutuhan pemesanan melalui QR Code meja dan pembayaran digital melalui Midtrans.",
    ),
    (
        "2. Identifikasi Masalah, Tujuan, dan Pengumpulan Data (paralel)",
        "Setelah analisis kebutuhan, penelitian memasuki tiga kegiatan yang digambarkan "
        "secara paralel pada Gambar 3.1. Pertama, identifikasi masalah difokuskan pada "
        "kendala proses pemesanan manual di Restoran Kekupu Villa Jembrana, misalnya "
        "antrian, kesalahan pencatatan pesanan, serta keterbatasan pelacakan status "
        "pesanan. Kedua, peneliti menentukan tujuan dan ruang lingkup penelitian agar "
        "sistem yang dibangun terbatas pada modul pemesanan berbasis QR Code, "
        "pengelolaan menu/pesanan, pembayaran Midtrans, serta laporan untuk admin. "
        "Ketiga, pengumpulan data dilakukan melalui observasi di lokasi restoran, "
        "wawancara dengan pihak pengelola/staf, dan studi pustaka terkait sistem "
        "informasi restoran, Laravel, QR Code, serta gateway pembayaran.",
    ),
    (
        "3. Desain Sistem (Use Case, ERD, DFD, UI)",
        "Hasil analisis dituangkan ke dalam desain sistem yang mencakup Use Case untuk "
        "aktor pelanggan, staf, dan admin; Entity Relationship Diagram (ERD) untuk "
        "struktur basis data pesanan, meja, menu, dan transaksi; Data Flow Diagram (DFD) "
        "untuk aliran data pemesanan dan pembayaran; serta desain antarmuka pengguna "
        "(UI) untuk halaman pelanggan berbasis QR dan panel admin Laravel.",
    ),
    (
        "4. Penulisan Kode",
        "Tahap implementasi dilakukan dengan menulis kode aplikasi menggunakan Laravel 12. "
        "Modul yang dikembangkan meliputi autentikasi admin/staf, pemesanan pelanggan "
        "melalui QR Code nomor meja, pengelolaan status pesanan, cetak nota, serta "
        "integrasi pembayaran QRIS melalui Midtrans sesuai kebutuhan operasional Restoran "
        "Kekupu Villa Jembrana.",
    ),
    (
        "5. Pengujian Sistem (Black-Box Testing & UAT)",
        "Pengujian sistem dilakukan dengan Black-Box Testing untuk memverifikasi "
        "kesesuaian fungsi input–output tanpa melihat struktur internal kode, serta "
        "User Acceptance Testing (UAT) bersama pihak Restoran Kekupu Villa untuk "
        "memastikan sistem diterima dan layak digunakan pada proses pemesanan nyata.",
    ),
    (
        "6. Penerapan dan Pemeliharaan",
        "Setelah pengujian dinyatakan memenuhi kebutuhan, sistem diterapkan pada "
        "lingkungan operasional Restoran Kekupu Villa Jembrana. Pemeliharaan mencakup "
        "perbaikan cacat minor, penyesuaian konfigurasi Midtrans, pembaruan data menu, "
        "serta monitoring kelancaran pemesanan berbasis QR Code.",
    ),
]

CLOSING = (
    "Dengan mengikuti tahapan pada Gambar 3.1, alur penelitian tidak lagi mengacu pada "
    "diagram IPO, melainkan pada siklus Waterfall yang runtut dari analisis kebutuhan "
    "hingga penerapan dan pemeliharaan. Tahap Selesai menandai berakhirnya prosedur "
    "penelitian setelah sistem dinyatakan siap digunakan."
)

CATATAN_PEMBIMBING = (
    "Catatan perbaikan untuk pembimbing: (a) ganti seluruh referensi gambar/alur IPO "
    "pada Bab III dengan Gambar 3.1 Waterfall; (b) pastikan narasi pengujian menyebut "
    "Black-Box Testing dan UAT secara eksplisit; (c) pastikan konteks objek penelitian "
    "adalah Restoran Kekupu Villa Jembrana (bukan contoh Ari Nata)."
)


def build_naskah() -> Path:
    doc = Document()
    setup_page(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Cm(0)
    r = title.add_run("NASKAH PERBAIKAN BAB III")
    _set_run_font(r, size=14, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.first_line_indent = Cm(0)
    r = sub.add_run(
        "Penggantian Alur IPO menjadi Alur Penelitian Metode Waterfall\n"
        "(Gambar 3.1) — Proposal Restoran Kekupu Villa Jembrana"
    )
    _set_run_font(r, size=12, bold=True)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.first_line_indent = Cm(0)
    meta.paragraph_format.space_after = Pt(12)
    r = meta.add_run(
        "Revisi dokumen proposal — September 2026\n"
        "Teknologi: Laravel 12 · QR Code meja · Midtrans · Black-Box Testing & UAT"
    )
    _set_run_font(r, size=11, italic=True)

    add_heading_tnr(doc, "1. Ringkasan Perbaikan", level=1)
    add_body(doc, INTRO_REPLACEMENT)
    add_body(doc, CATATAN_PEMBIMBING)

    add_heading_tnr(doc, "2. Gambar 3.1 yang Dimasukkan ke Naskah", level=1)
    add_body(
        doc,
        "Gambar berikut menggantikan diagram IPO pada bagian prosedur/alur penelitian. "
        "Silakan salin gambar beserta keterangan ke dalam file proposal utama.",
        first_line=True,
    )
    insert_gambar_3_1(doc)

    add_heading_tnr(doc, "3. Teks Pengganti untuk Subbab 3.4 Prosedur Penelitian", level=1)
    add_body(
        doc,
        "Paragraf di bawah ini dapat disalin langsung sebagai pengganti narasi prosedur "
        "penelitian (sebelumnya IPO) agar selaras dengan tahapan pada Gambar 3.1.",
        first_line=True,
    )
    add_body(doc, PROSEDUR_INTRO)

    for heading, para in STAGES:
        p = doc.add_paragraph()
        run = p.add_run(heading)
        _set_run_font(run, size=12, bold=True)
        _set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False, space_after=3)
        add_body(doc, para)

    add_body(doc, CLOSING)

    add_heading_tnr(doc, "4. Pemetaan Tahapan Gambar 3.1", level=1)
    add_body(
        doc,
        "Urutan resmi yang harus tercermin dalam naskah: Mulai → Analisis Kebutuhan → "
        "(paralel) Identifikasi Masalah | Menentukan Tujuan dan Ruang Lingkup Penelitian | "
        "Pengumpulan Data (observasi, wawancara, studi pustaka) → Desain Sistem "
        "(Use Case, ERD, DFD, UI) → Penulisan Kode → Pengujian Sistem "
        "(Black-Box Testing & UAT) → Penerapan dan Pemeliharaan → Selesai.",
        first_line=True,
    )

    out = OUT_DIR / "Naskah_Perbaikan_Bab3_Waterfall.docx"
    doc.save(out)
    return out


def build_proposal_revisi() -> Path:
    """Bab III excerpt ready to merge into the full proposal revision file."""
    doc = Document()
    setup_page(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Cm(0)
    r = title.add_run("PROPOSAL REVISI — 22 SEPTEMBER 2026")
    _set_run_font(r, size=14, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.first_line_indent = Cm(0)
    sub.paragraph_format.space_after = Pt(12)
    r = sub.add_run(
        "Kutipan Bab III (Metodologi Penelitian)\n"
        "Sistem Informasi Pemesanan Restoran Kekupu Villa Jembrana\n"
        "Berbasis Laravel 12, QR Code, dan Midtrans"
    )
    _set_run_font(r, size=12, bold=True)

    add_heading_tnr(doc, "BAB III METODE PENELITIAN", level=1)

    add_heading_tnr(doc, "3.1 Pendekatan Penelitian", level=2)
    add_body(
        doc,
        "Penelitian ini menggunakan pendekatan rekayasa perangkat lunak dengan model "
        "pengembangan Waterfall. Model tersebut dipilih karena kebutuhan sistem Restoran "
        "Kekupu Villa Jembrana telah dapat dirumuskan secara jelas pada tahap awal, "
        "sehingga pengembangan dapat dilakukan secara bertahap dan terdokumentasi.",
    )

    add_heading_tnr(doc, "3.2 Objek dan Lokasi Penelitian", level=2)
    add_body(
        doc,
        "Objek penelitian adalah proses pemesanan pada Restoran Kekupu Villa yang "
        "berlokasi di Dauhwaru, Pendem, Kecamatan Jembrana, Kabupaten Jembrana, Bali. "
        "Sistem yang dikembangkan mendukung pemesanan pelanggan melalui QR Code meja, "
        "pengelolaan pesanan oleh staf/admin, serta pembayaran melalui Midtrans.",
    )

    add_heading_tnr(doc, "3.3 Metode Pengumpulan Data", level=2)
    add_body(
        doc,
        "Pengumpulan data dilakukan melalui observasi kegiatan operasional restoran, "
        "wawancara dengan pengelola/staf terkait alur pemesanan dan pembayaran, serta "
        "studi pustaka mengenai sistem informasi restoran, kerangka kerja Laravel, "
        "teknologi QR Code, dan gateway pembayaran Midtrans.",
    )

    add_heading_tnr(doc, "3.4 Prosedur Penelitian", level=2)
    add_body(doc, INTRO_REPLACEMENT)
    add_body(doc, PROSEDUR_INTRO)
    insert_gambar_3_1(doc)

    for heading, para in STAGES:
        p = doc.add_paragraph()
        run = p.add_run(heading)
        _set_run_font(run, size=12, bold=True)
        _set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False, space_after=3)
        add_body(doc, para)

    add_body(doc, CLOSING)

    add_heading_tnr(doc, "3.5 Metode Pengujian", level=2)
    add_body(
        doc,
        "Sesuai Gambar 3.1, pengujian sistem mencakup Black-Box Testing untuk menguji "
        "fungsi-fungsi utama (pemindaian QR meja, pemesanan menu, status pesanan, "
        "pembayaran Midtrans, dan laporan admin) berdasarkan skenario input–output. "
        "Selanjutnya dilakukan User Acceptance Testing (UAT) bersama perwakilan Restoran "
        "Kekupu Villa untuk menilai kemudahan penggunaan dan kesesuaian dengan kebutuhan "
        "operasional sebelum penerapan penuh.",
    )

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(18)
    run = note.add_run(
        "Catatan: File ini merupakan naskah revisi Bab III yang diselaraskan dengan "
        "Gambar 3.1. Bagian Bab I–II dan Bab IV–V dapat digabung dari naskah proposal "
        "utama setelah diagram IPO diganti sepenuhnya dengan alur Waterfall ini."
    )
    _set_run_font(run, size=11, italic=True, color=RGBColor(0x33, 0x33, 0x33))
    _set_paragraph_format(note, first_line=False)

    out = OUT_DIR / "Proposal-Revisi-22-September-2026.docx"
    doc.save(out)
    return out


def copy_artifacts(*paths: Path) -> None:
    ARTIFACTS.mkdir(parents=True, exist_ok=True)
    for src in paths:
        dest = ARTIFACTS / src.name
        dest.write_bytes(src.read_bytes())
        print(f"Copied artifact: {dest}")


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    if not PNG.exists():
        raise SystemExit(f"PNG not found: {PNG}")

    naskah = build_naskah()
    proposal = build_proposal_revisi()
    print(f"Wrote: {naskah} ({naskah.stat().st_size} bytes)")
    print(f"Wrote: {proposal} ({proposal.stat().st_size} bytes)")
    copy_artifacts(naskah, proposal, PNG)


if __name__ == "__main__":
    main()
